"""Field records — the Epicollect download of a field trip and the fieldwork datasheet turned into the updated
datasheet, the deployments of the collection and its FileTimestampLog and FieldworkLog."""
import csv
import datetime as dt
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from fastapi.testclient import TestClient

from wildintel_uploader.core import config
from wildintel_uploader.core.services import fieldlog_service as fs
from wildintel_uploader.core.services import local_folder_service as lfs
from wildintel_uploader.web.main import app

HEADERS = ["Trappers", "Session", "Date", "Time", "DateTime", "Site", "Camera", "Action", "Status", "CardBatteriesOK",
           "NumPhotos", "SD", "DateTimeOK", "Notes", "Longitude", "Latitude"]
EPI_HEADERS = ["ec5_uuid", "created_at", "1_Responsable", "2_Collection", "3_Date", "4_Location", "5_Camera", "6_Action",
               "7_Status_of_the_came", "8_Card_Batteries", "9_Number_of_pictures", "10_SD", "11_Date_and_hour",
               "12_Observation", "lat_13_Coordenates", "long_13_Coordenates"]

# The datasheet: two sites installed in R0001 and revised in R0002.
SHEET_ROWS = [
    ("Ana", "R0001", dt.date(2026, 1, 10), dt.time(9, 0), "TATR_0001", 11, "Installation", "OK", None, "1a", "pole"),
    ("Ana", "R0001", dt.date(2026, 1, 10), dt.time(10, 0), "TATR_0002", 12, "Installation", "OK", None, "2a", None),
    ("Ben", "R0002", dt.date(2026, 2, 20), dt.time(9, 30), "TATR_0001", 11, "Revision", "OK", 120, "1b", None),
    ("Ben", "R0002", dt.date(2026, 2, 20), dt.time(11, 0), "TATR_0002", 12, "Revision", "OK", 80, "2b", "tilted"),
]
# The field trip of R0003: TATR_0001 revised (its card full: no count), TATR_0003 installed — a new site. TATR_0002 is
# not visited yet.
EPI_ROWS = [
    ("u1", "2026-03-30T08:15:42.000Z", "Carla", "3", "30/03/2026", "1", "11", "Revision", "OK", "OK", "", "1c", "OK", "card full", "49.2", "20.1"),
    ("u2", "2026-03-30T10:05:00.000Z", "Carla", "3", "30/03/2026", "3", "15", "Installation", "OK", "OK", "", "3a", "OK", "", "49.3", "20.2"),
]


def _datasheet(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for trappers, session, day, time, site, camera, action, status, photos, sd, notes in SHEET_ROWS:
        ws.append([trappers, session, day, time, None, site, camera, action, status, "OK", photos, sd, "OK", notes, 20.0, 49.0])
    wb.save(path)
    return path


def _epicollect(path: Path, rows=EPI_ROWS) -> Path:
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(EPI_HEADERS)
        writer.writerows(rows)
    return path


@pytest.fixture
def trip(tmp_path):
    """The two inputs, and a collections folder with the research project TATRA and its locations TATR_0001-0002."""
    root = tmp_path / "collections"
    lfs.write_research_project(root, {"acronym": "TATRA", "name": "Tatra"})
    for site in ("TATR_0001", "TATR_0002"):
        lfs.write_location(root, "TATRA", {"location_id": site, "timezone": "UTC"})
    sheet = _datasheet(tmp_path / "TatraFieldSheet_R0001-R0002.xlsx")
    epi = _epicollect(tmp_path / "form-backup.csv")
    with patch("wildintel_uploader.core.config.collections_dir", return_value=root):
        yield root, epi, sheet


def test_the_trip_closes_the_deployments_of_the_previous_collection(trip):
    _, epi, sheet = trip
    result = fs.analyze(str(epi), str(sheet), "TATRA")

    assert result["new_rows"] == 2
    assert result["collections"] == ["R0002"]
    by_id = {d["deployment_id"]: d for d in result["deployments"]}
    closed = by_id["R0002-TATR_0001"]
    assert closed["status"] == "closed"
    assert (closed["start"], closed["end"]) == ("2026-02-20T09:30:00Z", "2026-03-30T08:15:00Z")
    # The camera, the card and who set it up are the start visit's; a full card closes it all the same.
    assert (closed["setup_by"], closed["camera_id"], closed["card"], closed["photos_counted"]) == ("Ben", "11", "1b", None)
    assert closed["comments"] == "End: card full"
    # TATR_0002 is not visited yet: its camera is still out there.
    assert by_id["R0002-TATR_0002"]["status"] == "open"
    assert by_id["R0002-TATR_0002"]["comments"] == "Start: tilted"
    assert by_id["R0003-TATR_0003"]["status"] == "open"


def test_a_site_that_is_not_a_location_yet_is_reported(trip):
    _, epi, sheet = trip
    result = fs.analyze(str(epi), str(sheet), "TATRA")

    assert [s["site"] for s in result["new_sites"]] == ["TATR_0003"]
    assert result["new_sites"][0]["latitude"] == 49.3
    assert any("TATR_0003" in w and "Trapper" in w for w in result["warnings"])


def test_without_locations_kept_new_sites_are_not_judged(trip):
    _, epi, sheet = trip
    result = fs.analyze(str(epi), str(sheet), None)

    assert result["locations_known"] is False and result["new_sites"] == []


def test_a_camera_at_two_sites_in_a_session_is_reported(trip, tmp_path):
    _, _, sheet = trip
    rows = [EPI_ROWS[0], EPI_ROWS[1][:6] + ("11",) + EPI_ROWS[1][7:]]
    epi = _epicollect(tmp_path / "shared.csv", rows)

    result = fs.analyze(str(epi), str(sheet), "TATRA")

    assert any("camera 11 is at 2 sites" in w for w in result["warnings"])


def test_applying_writes_the_datasheet_and_the_collection_files(trip):
    root, epi, sheet = trip
    result = fs.apply(str(epi), str(sheet), "TATRA")

    updated = sheet.parent / "TatraFieldSheet_R0001-R0003.xlsx"
    assert str(updated) in result["files"] and sheet.is_file()
    assert openpyxl.load_workbook(updated).active.max_row == 7
    collection = root / "TATRA" / "R0002"
    log = lfs.read_timestamp_log(lfs.timestamp_log_path(collection))
    assert log == [{"Deployment": "R0002-TATR_0001", "StartDate": "2026:02:20", "StartTime": "09:30:00",
                    "EndDate": "2026:03:30", "EndTime": "08:15:00"}]
    rows = fs.read_fieldwork_log(collection)
    assert [(r["deploymentID"], r["status"]) for r in rows] == [("R0002-TATR_0001", "closed"), ("R0002-TATR_0002", "open")]
    assert rows[0]["setupBy"] == "Ben" and rows[0]["cameraID"] == "11" and rows[0]["deploymentEnd"] == "2026-03-30T08:15:00Z"


def test_applying_again_keeps_the_other_rows_of_the_timestamp_log(trip):
    root, epi, sheet = trip
    collection = root / "TATRA" / "R0002"
    lfs.upsert_timestamp_log(collection, "R0002-TATR_0099", dt.datetime(2026, 2, 1), dt.datetime(2026, 3, 1))

    fs.apply(str(epi), str(sheet), "TATRA")

    ids = [r["Deployment"] for r in lfs.read_timestamp_log(lfs.timestamp_log_path(collection))]
    assert ids == ["R0002-TATR_0099", "R0002-TATR_0001"]


def test_a_research_project_is_needed_to_write(trip):
    _, epi, sheet = trip
    with pytest.raises(fs.FieldlogError):
        fs.apply(str(epi), str(sheet), "")
    with pytest.raises(fs.FieldlogError):
        fs.apply(str(epi), str(sheet), "NOPE")


def test_a_missing_input_is_an_error(trip, tmp_path):
    _, epi, _ = trip
    with pytest.raises(fs.FieldlogError, match="datasheet"):
        fs.analyze(str(epi), str(tmp_path / "missing.xlsx"))


def test_columns_named_otherwise_are_mapped_in_the_settings(trip, tmp_path):
    _, _, sheet = trip
    renamed = _epicollect(tmp_path / "renamed.csv")
    text = renamed.read_text(encoding="utf-8").replace("12_Observation", "12_Remarks", 1)
    renamed.write_text(text, encoding="utf-8")
    settings = config.Settings()
    with pytest.raises(fs.FieldlogError, match="Observation"):
        fs.analyze(str(renamed), str(sheet), "TATRA", settings)

    settings.FIELDLOG.columns = {"Notes": "12_Remarks"}
    result = fs.analyze(str(renamed), str(sheet), "TATRA", settings)

    assert result["new_rows"] == 2


def test_the_api_analyzes_and_applies(trip):
    _, epi, sheet = trip
    client = TestClient(app)
    payload = {"epicollect": str(epi), "datasheet": str(sheet), "research_project_id": "TATRA"}

    analyzed = client.post("/api/fieldlog/analyze", json=payload)
    applied = client.post("/api/fieldlog/apply", json=payload)
    last = client.get("/api/fieldlog/last").json()

    assert analyzed.status_code == 200 and analyzed.json()["collections"] == ["R0002"]
    assert applied.status_code == 200
    assert last["research_project_id"] == "TATRA" and last["datasheet"].endswith("R0001-R0003.xlsx")
    assert client.post("/api/fieldlog/analyze", json={**payload, "datasheet": "/nope.xlsx"}).status_code == 400
    assert len(client.get("/api/fieldlog/columns").json()["results"]) == 16


def test_the_camera_timezone_must_be_a_timezone():
    with pytest.raises(ValueError):
        config.FieldlogSettings(camera_timezone="Mars/Olympus")
    assert config.FieldlogSettings().camera_timezone == "UTC"
